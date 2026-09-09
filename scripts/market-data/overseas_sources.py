#!/usr/bin/env python3
"""Trusted overseas directory sources and isolated raw SINA candle worker."""

import argparse
import csv
import io
import json
import math
import os
import re
import subprocess
import sys
import tempfile
import threading
from collections import OrderedDict
from datetime import datetime, time, timedelta, timezone
from decimal import Decimal, InvalidOperation
from zoneinfo import ZoneInfo


HKEX_ENGLISH_URL = "https://www.hkex.com.hk/eng/services/trading/securities/securitieslists/ListOfSecurities.xlsx"
HKEX_CHINESE_URL = "https://www.hkex.com.hk/chi/services/trading/securities/securitieslists/ListOfSecurities_c.xlsx"
NASDAQ_LISTED_URL = "https://www.nasdaqtrader.com/dynamic/SymDir/nasdaqlisted.txt"
NASDAQ_OTHER_URL = "https://www.nasdaqtrader.com/dynamic/SymDir/otherlisted.txt"

MAX_DIRECTORY_BYTES = 10_000_000
MAX_WORKER_OUTPUT_BYTES = 20_000_000
MAX_DIRECTORY_ITEMS = 30_000
MAX_WORKBOOK_ROWS = 30_000

HK_CODE = re.compile(r"^[0-9]{1,5}$")
US_CODE = re.compile(r"^[A-Z][A-Z0-9.-]{0,9}$")
US_PREFERRED_CODE = re.compile(r"^[A-Z][A-Z0-9.-]{0,8}\$[A-Z]$")
US_NON_EQUITY_NAME = re.compile(
    r"\b(?:warrants?|units?|rights?|preferred|preference|pfd|notes?|bonds?|debentures?|debt)\b",
    re.IGNORECASE,
)
OTHER_EXCHANGES = {
    "A": "NYSE AMERICAN",
    "N": "NYSE",
    "P": "NYSE ARCA",
    "Z": "CBOE BZX",
    "V": "IEX",
}


def _download_bytes(url, max_bytes=MAX_DIRECTORY_BYTES, connect_timeout=5, read_timeout=70):
    import requests

    response = requests.get(url, stream=True, timeout=(connect_timeout, read_timeout))
    response.raise_for_status()
    content_length = response.headers.get("Content-Length")
    if content_length is not None and int(content_length) > max_bytes:
        raise ValueError("upstream response is too large")
    chunks = []
    size = 0
    for chunk in response.iter_content(chunk_size=64 * 1024):
        if not chunk:
            continue
        size += len(chunk)
        if size > max_bytes:
            raise ValueError("upstream response is too large")
        chunks.append(chunk)
    return b"".join(chunks)


def _workbook_rows(payload, required_headers):
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        # HKEX currently declares A1:R8 despite containing more than 17k rows.
        # Read-only openpyxl trusts that stale dimension unless it is reset.
        sheet.reset_dimensions()
        return workbook_values_to_rows(sheet.iter_rows(values_only=True), required_headers)
    finally:
        workbook.close()


def workbook_values_to_rows(values, required_headers):
    iterator = iter(values)
    next(iterator, None)  # title
    next(iterator, None)  # update timestamp
    raw_headers = next(iterator, None)
    if raw_headers is None:
        raise ValueError("workbook is missing header row 3")
    headers = ["" if value is None else str(value).strip() for value in raw_headers]
    if not set(required_headers).issubset(headers):
        raise ValueError("workbook has unexpected headers")
    rows = []
    for row_count, row_values in enumerate(iterator, start=1):
        if row_count > MAX_WORKBOOK_ROWS:
            raise ValueError("workbook has too many rows")
        if any(value is not None for value in row_values):
            rows.append({header: value for header, value in zip(headers, row_values) if header})
    return rows


def _hk_symbol(raw):
    if isinstance(raw, bool):
        raise ValueError("invalid HK stock code")
    if isinstance(raw, float):
        if not raw.is_integer():
            raise ValueError("invalid HK stock code")
        raw = int(raw)
    value = str(raw).strip()
    if value.endswith(".0") and value[:-2].isdigit():
        value = value[:-2]
    if not HK_CODE.fullmatch(value) or int(value) <= 0:
        raise ValueError("invalid HK stock code")
    return value.zfill(5)


def build_hk_directory(english_rows, chinese_rows=None, min_items=1000):
    chinese_names = {}
    for row in chinese_rows or []:
        try:
            symbol = _hk_symbol(row.get("股份代號"))
            name = str(row.get("股份名稱") or "").strip()
        except (TypeError, ValueError):
            continue
        if name and len(name) <= 200:
            chinese_names[symbol] = name

    by_symbol = {}
    for row in english_rows:
        if str(row.get("Category") or "").strip().casefold() != "equity":
            continue
        symbol = _hk_symbol(row.get("Stock Code"))
        english_name = str(row.get("Name of Securities") or "").strip()
        currency = str(row.get("Trading Currency") or "").strip().upper()
        if not english_name or len(english_name) > 200 or not re.fullmatch(r"[A-Z]{3}", currency):
            raise ValueError("invalid HK equity metadata")
        if symbol in by_symbol:
            raise ValueError("duplicate HK equity")
        by_symbol[symbol] = {
            "symbol": symbol,
            "name": chinese_names.get(symbol, english_name),
            "search_names": list(dict.fromkeys([english_name, chinese_names.get(symbol, english_name)])),
            "market": "HK",
            "currency": currency,
            "exchange": "HKEX",
            "timezone": "Asia/Hong_Kong",
        }
    if not min_items <= len(by_symbol) <= MAX_DIRECTORY_ITEMS:
        raise ValueError("HK directory has an implausible size")
    return [by_symbol[symbol] for symbol in sorted(by_symbol)]


def load_hk_directory(download=_download_bytes, min_items=1000):
    english_rows = _workbook_rows(
        download(HKEX_ENGLISH_URL),
        ("Stock Code", "Name of Securities", "Category", "Trading Currency"),
    )
    chinese_rows = []
    try:
        chinese_rows = _workbook_rows(
            download(HKEX_CHINESE_URL),
            ("股份代號", "股份名稱", "分類", "交易貨幣"),
        )
    except Exception:
        # Chinese names are enrichment only; English identity and currency stay authoritative.
        chinese_rows = []
    return build_hk_directory(english_rows, chinese_rows, min_items=min_items)


def _is_named_equity(name):
    return bool(name) and len(name) <= 200 and not US_NON_EQUITY_NAME.search(name)


def parse_nasdaq_file(payload, listed):
    try:
        text = payload.decode("ascii")
    except (AttributeError, UnicodeDecodeError) as exception:
        raise ValueError("Nasdaq directory is not ASCII") from exception
    lines = [line.rstrip("\r") for line in text.splitlines() if line.strip()]
    listed_header = [
        "Symbol", "Security Name", "Market Category", "Test Issue", "Financial Status",
        "Round Lot Size", "ETF", "NextShares",
    ]
    other_header = [
        "ACT Symbol", "Security Name", "Exchange", "CQS Symbol", "ETF",
        "Round Lot Size", "Test Issue", "NASDAQ Symbol",
    ]
    expected_header = listed_header if listed else other_header
    if len(lines) < 3 or lines[0].split("|") != expected_header:
        raise ValueError("Nasdaq directory has unexpected headers")
    if not lines[-1].startswith("File Creation Time:"):
        raise ValueError("Nasdaq directory is missing completion footer")

    items = []
    reader = csv.DictReader(lines[1:-1], fieldnames=expected_header, delimiter="|")
    for row in reader:
        if None in row or row["Test Issue"] != "N" or row["ETF"] != "N":
            if None in row:
                raise ValueError("Nasdaq directory row has unexpected fields")
            continue
        symbol = row["Symbol" if listed else "ACT Symbol"].strip().upper()
        name = row["Security Name"].strip()
        if not _is_named_equity(name) or US_PREFERRED_CODE.fullmatch(symbol):
            continue
        if not US_CODE.fullmatch(symbol):
            raise ValueError("invalid US symbol")
        if listed:
            if row["Market Category"] not in {"Q", "G", "S"}:
                raise ValueError("invalid Nasdaq market category")
            exchange = "NASDAQ"
        else:
            exchange_code = row["Exchange"].strip().upper()
            if exchange_code not in OTHER_EXCHANGES:
                raise ValueError("invalid US exchange code")
            exchange = OTHER_EXCHANGES[exchange_code]
        items.append({
            "symbol": symbol,
            "name": name,
            "market": "US",
            "currency": "USD",
            "exchange": exchange,
            "timezone": "America/New_York",
        })
    return items


def load_us_directory(download=_download_bytes, min_items=2000):
    # Fetch both complete files before parsing or publishing either one.
    listed_payload = download(NASDAQ_LISTED_URL)
    other_payload = download(NASDAQ_OTHER_URL)
    combined = parse_nasdaq_file(listed_payload, listed=True) + parse_nasdaq_file(other_payload, listed=False)
    by_symbol = {}
    for item in combined:
        if item["symbol"] in by_symbol:
            raise ValueError("duplicate US symbol")
        by_symbol[item["symbol"]] = item
    if not min_items <= len(by_symbol) <= MAX_DIRECTORY_ITEMS:
        raise ValueError("US directory has an implausible size")
    return [by_symbol[symbol] for symbol in sorted(by_symbol)]


def load_directory(market):
    if market == "HK":
        return load_hk_directory()
    if market == "US":
        return load_us_directory()
    raise ValueError("market must be HK or US")


def _set_safe_v8_flags():
    from py_mini_racer import MiniRacer

    # The Linux dependency used by AKShare reads this class value on first context creation.
    MiniRacer.v8_flags = ["--single-threaded", "--jitless"]
    return MiniRacer


def _plain_number(value):
    if value is None:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def normalize_sina_records(records, market):
    rows = []
    for raw in records:
        raw_date = raw.get("date")
        if hasattr(raw_date, "isoformat"):
            raw_date = raw_date.isoformat()
        rows.append({
            "date": str(raw_date),
            "open": _plain_number(raw.get("open")),
            "high": _plain_number(raw.get("high")),
            "low": _plain_number(raw.get("low")),
            "close": _plain_number(raw.get("close")),
            "volume": _plain_number(raw.get("volume")),
            # SINA's documented US raw contract is OHLCV only. Its decoded amount
            # field is not verified turnover and must remain unavailable.
            "turnover": None if market == "US" else _plain_number(raw.get("amount", raw.get("turnover"))),
        })
    return rows


def load_us_sina_rows(symbol, download=_download_bytes, decode=None):
    symbol = str(symbol).strip().upper()
    if not US_CODE.fullmatch(symbol):
        raise ValueError("invalid US symbol")
    payload = download(f"https://finance.sina.com.cn/staticdata/us/{symbol}")
    text = payload.decode("utf-8")
    if "=" not in text or ";" not in text:
        raise ValueError("SINA history has unexpected framing")
    encoded = text.split("=", 1)[1].split(";", 1)[0].strip().strip('"')
    if not encoded or len(encoded) > MAX_WORKER_OUTPUT_BYTES:
        raise ValueError("SINA history payload is invalid")
    if decode is None:
        MiniRacer = _set_safe_v8_flags()
        from akshare.stock.cons import zh_js_decode

        context = MiniRacer()
        context.eval(zh_js_decode)  # Installed AKShare decoder constant, never remote source.

        def decode(value):
            return context.call("d", value)
    return normalize_sina_records(decode(encoded), "US")


def load_sina_rows(market, symbol):
    if market == "HK":
        _set_safe_v8_flags()
        import akshare as ak

        # Raw adjustment is the only stock_hk_daily branch that does not fetch/eval factors.
        frame = ak.stock_hk_daily(symbol=symbol, adjust="")
        return normalize_sina_records(frame.to_dict("records"), market)
    if market == "US":
        return load_us_sina_rows(symbol)
    raise ValueError("market must be HK or US")


def run_candle_worker(market, symbol, timeout_seconds=20):
    command = [sys.executable, os.path.abspath(__file__), "--worker", market, symbol]
    environment = {
        "PATH": "/usr/bin:/bin:/usr/local/bin",
        "PYTHONNOUSERSITE": "1",
        "PYTHONUNBUFFERED": "1",
    }
    with tempfile.TemporaryFile() as output:
        process = subprocess.Popen(
            command,
            stdin=subprocess.DEVNULL,
            stdout=output,
            stderr=subprocess.DEVNULL,
            env=environment,
        )
        try:
            process.wait(timeout=timeout_seconds)
        except subprocess.TimeoutExpired as exception:
            process.kill()
            process.wait()
            raise RuntimeError("SINA worker timed out") from exception
        output.seek(0, os.SEEK_END)
        size = output.tell()
        if process.returncode != 0 or size > MAX_WORKER_OUTPUT_BYTES:
            raise RuntimeError("SINA worker failed")
        output.seek(0)
        try:
            payload = json.loads(output.read().decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exception:
            raise RuntimeError("SINA worker returned invalid JSON") from exception
    if not isinstance(payload, dict) or not isinstance(payload.get("rows"), list):
        raise RuntimeError("SINA worker returned invalid payload")
    return payload["rows"]


SPOT_ZONES = {'CN': 'Asia/Shanghai', 'HK': 'Asia/Hong_Kong', 'US': 'America/New_York'}
SPOT_CURRENCIES = {'CN': 'CNY', 'HK': 'HKD', 'US': 'USD'}


def spot_symbols(market, symbols):
    if market not in SPOT_ZONES or not isinstance(symbols, list) or not 1 <= len(symbols) <= 50:
        raise ValueError('expected CN/HK/US and 1-50 symbols')
    pattern = r'[0-9]{6}\.(SH|SZ|BJ)' if market == 'CN' else r'[0-9]{5}' if market == 'HK' else r'[A-Z][A-Z0-9.-]{0,9}'
    if any(not isinstance(s, str) or not re.fullmatch(pattern, s) for s in symbols):
        raise ValueError('invalid spot symbol')
    return list(dict.fromkeys(symbols))


def spot_key(market, symbol):
    return symbol[-2:].lower() + symbol[:6] if market == 'CN' else ('hk' if market == 'HK' else 'us') + symbol


def completed_market_day(now, market):
    local = now.astimezone(ZoneInfo(SPOT_ZONES[market]))
    # Conservative publication buffer; provider must still return a real bar for this date.
    day = local.date() if local.time() >= time(17, 30) else local.date() - timedelta(days=1)
    while day.weekday() >= 5:
        day -= timedelta(days=1)
    return day


def candle_cache_current(fetched, as_of, now, market, ttl):
    target=completed_market_day(now,market)
    budget=min(ttl,900) if not as_of or as_of < target.isoformat() else ttl
    return (now-fetched).total_seconds()<budget and completed_market_day(fetched,market)==target


def before_latest_close(stamp, now, market):
    zone=ZoneInfo(SPOT_ZONES[market]);local=now.astimezone(zone);quote=stamp.astimezone(zone)
    close=time(15) if market=='CN' else time(16)
    day=local.date() if local.time()>=close else local.date()-timedelta(days=1)
    while day.weekday()>=5: day-=timedelta(days=1)
    return quote.date()<day or (quote.date()==day and quote.time()<close)


def spot_session(now, market):
    local = now.astimezone(ZoneInfo(SPOT_ZONES[market]))
    windows = [(time(9, 30), time(16, 30))] if market == 'US' else ([(time(9, 30), time(11, 30)), (time(13), time(15, 30))] if market == 'CN' else [(time(9, 30), time(12)), (time(13), time(16, 30))])
    if local.weekday() < 5 and any(start <= local.time() < end for start, end in windows):
        return 'TRADING_HOURS', 60
    for offset in range(8):
        day = local.date() + timedelta(days=offset)
        if day.weekday() >= 5:
            continue
        for start, _ in windows:
            candidate = datetime.combine(day, start, tzinfo=local.tzinfo)
            if candidate > local:
                return 'CLOSED_HOURS', max(60, int((candidate.astimezone(timezone.utc) - now).total_seconds()))
    return 'CLOSED_HOURS', 86400


def fetch_spot_text(market, symbols):
    import requests
    symbols = spot_symbols(market, symbols)
    url = 'https://qt.gtimg.cn/q=' + ','.join(spot_key(market, s) for s in symbols)
    with requests.get(url, timeout=(3, 7), stream=True, allow_redirects=False) as response:
        if response.status_code != 200:
            raise RuntimeError('spot provider unavailable')
        payload = bytearray()
        for chunk in response.iter_content(8192):
            payload.extend(chunk)
            if len(payload) > 262144:
                raise ValueError('spot response too large')
    return payload.decode('gb18030')


def parse_spot(raw, market, symbols, now):
    symbols = spot_symbols(market, symbols)
    if not isinstance(raw, str) or len(raw) > 262144 or now.tzinfo is None:
        raise ValueError('invalid spot response')
    expected = {spot_key(market, s): s for s in symbols}
    result, seen = {}, set()
    for key, body in re.findall(r'v_([A-Za-z0-9_.-]+)="([^"\r\n]*)";?', raw):
        if key not in expected:
            continue
        symbol = expected[key]
        if symbol in seen:
            raise ValueError('duplicate spot identity')
        seen.add(symbol)
        fields = body.split('~')
        try:
            expected_code = symbol[:6] if market == 'CN' else symbol
            allowed_codes = {expected_code} if market != 'US' else {expected_code, expected_code+'.OQ', expected_code+'.N', expected_code+'.A'}
            marker={'SH':'1','SZ':'51','BJ':'62'}[symbol[-2:]] if market=='CN' else {'HK':'100','US':'200'}[market]
            if len(fields) < 36 or fields[2] not in allowed_codes or fields[0] != marker or SPOT_CURRENCIES[market] not in fields:
                continue
            price = Decimal(fields[3])
            if not price.is_finite() or not 0 < price <= Decimal('1000000000000') or price.as_tuple().exponent < -6:
                continue
            fmt = '%Y%m%d%H%M%S' if market == 'CN' else '%Y/%m/%d %H:%M:%S' if market == 'HK' else '%Y-%m-%d %H:%M:%S'
            stamp = datetime.strptime(fields[30], fmt).replace(tzinfo=ZoneInfo(SPOT_ZONES[market])).astimezone(timezone.utc)
            if stamp > now + timedelta(seconds=120) or now - stamp > timedelta(days=7):
                continue
            result[symbol] = {'symbol':symbol,'market':market,'currency':SPOT_CURRENCIES[market], 'price':format(price,'f'),
                              'quotedAt':stamp.isoformat().replace('+00:00','Z'),'fetchedAt':now.isoformat().replace('+00:00','Z'),
                              'source':'TENCENT_PUBLIC','delayMinutes':None}
        except (ValueError, InvalidOperation, OverflowError):
            continue
    return result


class SpotQuoteService:
    """Read-only public quotes; shared bounded cache, no account/household data."""
    def __init__(self, loader, clock=None):
        self._loader = loader
        self._clock = clock or (lambda: datetime.now(timezone.utc))
        self._cache, self._attempts = OrderedDict(), OrderedDict()
        self._lock = threading.Lock()
        self._market_locks = {m:threading.Lock() for m in SPOT_ZONES}

    def quotes(self, market, symbols):
        symbols = spot_symbols(market, symbols)
        now = self._clock()
        state, interval = spot_session(now, market)
        due = []
        with self._lock:
            for symbol in symbols:
                key = (market, symbol)
                cached = self._cache.get(key)
                attempt = self._attempts.get(key)
                if (not cached or cached[1] <= now) and (not attempt or (now-attempt).total_seconds() >= 60):
                    due.append(symbol)
        if due and self._market_locks[market].acquire(timeout=0.2):
            try:
                with self._lock:
                    due = [s for s in due if not self._attempts.get((market,s)) or (now-self._attempts[(market,s)]).total_seconds() >= 60]
                    for symbol in due:
                        self._attempts[(market,symbol)] = now
                    while len(self._attempts) > 4096: self._attempts.popitem(last=False)
                if due:
                    try:
                        received = parse_spot(self._loader(market,due),market,due,self._clock())
                    except Exception:
                        received = {}
                    with self._lock:
                        for symbol in due:
                            key=(market,symbol)
                            if symbol in received:
                                stamp=datetime.fromisoformat(received[symbol]['quotedAt'].replace('Z','+00:00'))
                                expiry=min(interval,900) if before_latest_close(stamp,now,market) else interval
                                self._cache[key] = (received[symbol], now+timedelta(seconds=expiry), False)
                                self._cache.move_to_end(key)
                            elif key in self._cache:
                                previous=self._cache[key]
                                self._cache[key]=(previous[0],now+timedelta(seconds=60),True)
                        while len(self._cache) > 2048: self._cache.popitem(last=False)
            finally:
                self._market_locks[market].release()
        quotes=[]
        with self._lock:
            for symbol in symbols:
                cached=self._cache.get((market,symbol))
                quote=dict(cached[0]) if cached else {'symbol':symbol,'market':market,'currency':SPOT_CURRENCIES[market], 'price':None,'quotedAt':None,'fetchedAt':None,'source':'TENCENT_PUBLIC','delayMinutes':None}
                age = max(0,int((now-datetime.fromisoformat(quote['quotedAt'].replace('Z','+00:00'))).total_seconds())) if quote['quotedAt'] else None
                if age is not None and age > 604800: quote['price']=None
                old_close=quote['quotedAt'] and before_latest_close(datetime.fromisoformat(quote['quotedAt'].replace('Z','+00:00')),now,market)
                quote['status']='UNAVAILABLE' if quote['price'] is None else 'STALE' if cached[2] or old_close else 'DELAYED' if state=='TRADING_HOURS' and age>180 else 'OK'
                quote['ageSeconds']=age
                quotes.append(quote)
        # Regular hours are not a holiday calendar. Old source timestamps remain explicit.
        if any(q['status']=='UNAVAILABLE' for q in quotes): interval=min(interval,60)
        elif any(q['status']=='STALE' for q in quotes): interval=min(interval,900)
        elif state=='TRADING_HOURS' and quotes and all(q['ageSeconds'] is None or q['ageSeconds']>21600 for q in quotes): interval=300
        return {'quotes':quotes,'marketState':state,'nextRefreshSeconds':interval}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--worker", action="store_true")
    parser.add_argument("market", nargs="?")
    parser.add_argument("symbol", nargs="?")
    arguments = parser.parse_args()
    if not arguments.worker or arguments.market not in {"HK", "US"} or not arguments.symbol:
        raise SystemExit(2)
    print(json.dumps(
        {"rows": load_sina_rows(arguments.market, arguments.symbol)},
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    ))


if __name__ == "__main__":
    main()
